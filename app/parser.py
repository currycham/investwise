"""把用户粘贴的文本或上传的文件解析成标准化的公司列表。"""
import csv
import io
import re

from . import llm
from .prompts import PARSE_SYSTEM

MAX_TEXT = 12000


def file_to_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    lines.append("\t".join(cells))
        return "\n".join(lines)
    text = _decode(data)
    if name.endswith(".csv"):
        return "\n".join("\t".join(r) for r in csv.reader(io.StringIO(text)) if any(r))
    return text


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def normalize(ticker: str, market: str | None = None) -> tuple[str, str] | None:
    """返回 (market, ticker)；无法识别返回 None。"""
    t = ticker.strip().upper()
    t = re.sub(r"^(SH|SZ|BJ)(?=\d{6}$)", "", t)
    m = re.fullmatch(r"(\d{6})(?:\.(SH|SZ|SS|BJ))?", t)
    if m and market in (None, "A"):
        return "A", m.group(1)
    m = re.fullmatch(r"(\d{1,5})(?:\.HK)?", t)
    if m and market in (None, "HK"):
        return "HK", m.group(1).zfill(5)
    m = re.fullmatch(r"([A-Z]{1,5}(?:[.-][A-Z])?)(?:\.US)?", t)
    if m and market in (None, "US"):
        return "US", m.group(1).replace("-", ".")
    return None


def parse_companies(text: str, user_id: int | None = None) -> dict:
    text = text.strip()[:MAX_TEXT]
    if not text:
        return {"companies": [], "uncertain": []}
    try:
        result = llm.chat_json(
            [{"role": "system", "content": PARSE_SYSTEM}, {"role": "user", "content": text}],
            temperature=0, task="parse", user_id=user_id,
        )
    except llm.LLMError:
        return _regex_fallback(text)

    companies, seen = [], set()
    uncertain = list(result.get("uncertain") or [])
    for c in result.get("companies") or []:
        norm = normalize(str(c.get("ticker", "")), c.get("market"))
        if not norm:
            uncertain.append({"raw": f"{c.get('name')} {c.get('ticker')}", "reason": "代码格式无法识别"})
            continue
        if norm in seen:
            continue
        seen.add(norm)
        companies.append({"market": norm[0], "ticker": norm[1],
                          "name": (c.get("name") or norm[1]).strip(), "note": c.get("note") or ""})
    return {"companies": companies, "uncertain": uncertain}


def _regex_fallback(text: str) -> dict:
    """没有配置模型时，只识别显式写出的代码。"""
    companies, seen = [], set()
    for token in re.findall(r"\b(?:SH|SZ)?\d{6}(?:\.(?:SH|SZ|SS|BJ))?\b|\b\d{4,5}\.HK\b", text, re.I):
        norm = normalize(token)
        if norm and norm not in seen:
            seen.add(norm)
            companies.append({"market": norm[0], "ticker": norm[1], "name": norm[1],
                              "note": "未配置模型，名称请手动补充"})
    return {"companies": companies, "uncertain": []}
